"""Apply a blueprint patch (JSON, drafted by /decision-draft) to the workbook, then export it. The owner runs it in
their own terminal; Claude only drafts the patch (the guard asks before a shell command runs this tool).
Usage: python tools/blueprint_apply.py <patch.json> [--workbook docs/product-domain-blueprint-MVP.xlsx] [--out docs/blueprint]

The patch is an object with two lists, both optional:
  "rows":  {"sheet", "after": "<ID of the row above>" or "end", "values": {"<column header>": "<text>", ...}}
  "cells": {"sheet", "id": "<ID in the first column>", "column": "<header>", "old": "<current text>", "new": "<text>"}
"after" may name a row added earlier in the same patch; two rows of a sheet never share one "after" (to add rows
one under the other, name the row added just before). A newline in a value is a line break in the cell.

Every operation is checked before anything is written: the sheet and the columns exist, a new ID is not there yet,
the row "after" exists, "old" is exactly the current text (the patch was drafted from this workbook), no cell
changes an ID, and every value is text without a hidden character that does not start with = (Excel would read a
formula). New rows take the style and the height of the row above, and their values are written as text. The
autofilter (and any Excel table) of a changed sheet is extended to its last row.

The changed workbook is saved in a temporary folder next to it (.blueprint_apply-*) and exported there first, so
the ID checks of export_blueprint.py run too; only when they pass does it replace the workbook and copy the
Markdown files into the output folder. On any problem it lists every one, writes nothing and exits with 1.

Known limit: openpyxl's insert_rows does not move merged cells, data validation or conditional formatting ranges.
The blueprint sheets are flat tables; open the workbook in Excel once to look at the new rows."""
import argparse
import json
import os
import shutil
import sys
import tempfile
import zipfile
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.exceptions import InvalidFileException

sys.path.insert(0, str(Path(__file__).resolve().parent))
import export_blueprint  # noqa: E402

REPO = Path(__file__).resolve().parents[1]
WORKBOOK = REPO / 'docs' / 'product-domain-blueprint-MVP.xlsx'
OUT = REPO / 'docs' / 'blueprint'


def headers(ws):
    return [c.value for c in ws[1]]


def row_of(ws, rid):
    """Row number of the ID in the first column, or None."""
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == rid:
            return r
    return None


def last_row(ws):
    r = ws.max_row
    while r > 1 and all(ws.cell(r, c).value in (None, '') for c in range(1, ws.max_column + 1)):
        r -= 1
    return r


def text(v):
    """A value as the text the patch compares and writes: no carriage return, as in the export."""
    return '' if v is None else str(v).replace('\r\n', '\n').replace('\r', '')


def value_problems(where, name, v):
    if not isinstance(v, str):
        return [f'{where}: {name} is not text']
    found = []
    if v.startswith('='):
        found.append(f'{where}: {name} starts with =, which Excel reads as a formula')
    hidden = export_blueprint.hidden(text(v))
    if hidden:
        found.append(f'{where}: {name} has the hidden character(s) {" ".join(hidden)}')
    return found


def objects(patch, key, errors):
    ops = patch.get(key, [])
    if not isinstance(ops, list) or not all(isinstance(op, dict) for op in ops):
        errors.append(f'"{key}" is not a list of objects')
        return []
    return ops


def check(wb, patch):
    """Every problem of the patch against this workbook; empty when it can be applied."""
    if not isinstance(patch, dict):
        return ['the patch is not a JSON object with "rows" and "cells"']
    errors = [f'unknown key {k!r} (only "rows" and "cells")' for k in patch if k not in ('rows', 'cells')]
    rows, cells = objects(patch, 'rows', errors), objects(patch, 'cells', errors)
    changed = set()
    for i, op in enumerate(cells):
        where = f'cells[{i}]'
        sheet = op.get('sheet')
        if sheet not in wb.sheetnames:
            errors.append(f'{where}: no sheet {sheet!r}')
            continue
        ws = wb[sheet]
        head = headers(ws)
        column = op.get('column')
        if not isinstance(column, str) or column not in head:
            errors.append(f'{where}: no column {column!r} in {sheet}')
            continue
        if column == head[0]:
            errors.append(f'{where}: {column} is the ID column; an ID is not changed by a patch')
            continue
        r = row_of(ws, op.get('id'))
        if not r:
            errors.append(f'{where}: {op.get("id")!r} is not in {sheet}')
            continue
        if (sheet, r, column) in changed:
            errors.append(f'{where}: {op["id"]} {column} is changed twice in the patch')
        changed.add((sheet, r, column))
        current = text(ws.cell(r, head.index(column) + 1).value)
        if not isinstance(op.get('old'), str) or current != text(op['old']):
            errors.append(f'{where}: {op["id"]} {column} is {current!r}, the patch expects {op.get("old")!r}')
        errors += value_problems(where, 'the new text', op.get('new'))
    added, afters = {}, {}
    for i, op in enumerate(rows):
        where = f'rows[{i}]'
        sheet = op.get('sheet')
        if sheet not in wb.sheetnames:
            errors.append(f'{where}: no sheet {sheet!r}')
            continue
        ws = wb[sheet]
        head = headers(ws)
        planned = added.setdefault(sheet, set())
        anchors = afters.setdefault(sheet, set())
        values = op.get('values')
        if not isinstance(values, dict):
            errors.append(f'{where}: "values" is not an object')
            continue
        unknown = [k for k in values if k not in head]
        if unknown:
            errors.append(f'{where}: no column {", ".join(map(repr, unknown))} in {sheet}')
        rid = values.get(head[0])
        if not rid:
            errors.append(f'{where}: no value for the ID column {head[0]!r}')
        elif isinstance(rid, str) and (rid in planned or row_of(ws, rid)):
            errors.append(f'{where}: {rid} is already in {sheet}')
        after = op.get('after')
        if after != 'end' and not (isinstance(after, str) and (after in planned or row_of(ws, after))):
            errors.append(f'{where}: the row after which to insert, {after!r}, is not in {sheet}')
        elif after != 'end' and after in anchors:
            errors.append(f'{where}: another row is already inserted after {after}; name the row added before it')
        if isinstance(after, str):
            anchors.add(after)
        for k, v in values.items():
            errors += value_problems(where, k, v)
        if isinstance(rid, str):
            planned.add(rid)
    return errors


def insert(ws, after, values):
    head = headers(ws)
    above = last_row(ws) if after == 'end' else row_of(ws, after)
    r = above + 1
    # insert_rows moves the cells but not the row heights: move those of the rows below by hand.
    dims = ws.row_dimensions
    below = {k: dims[k].height for k in list(dims.keys()) if k >= r}
    ws.insert_rows(r)
    for k in below:
        dims[k].height = None
    for k, height in below.items():
        dims[k + 1].height = height
    dims[r].height = dims[above].height
    for c in range(1, len(head) + 1):
        src, dst = ws.cell(above, c), ws.cell(r, c)
        if src.has_style:
            dst._style = copy(src._style)
        v = values.get(head[c - 1])
        if v not in (None, ''):
            dst.value = text(v)
            dst.data_type = 's'


def extend(ref, last):
    """The range ref with its last row moved down to last (never up)."""
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    return f'{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max(max_row, last)}'


def apply(wb, patch):
    for op in patch.get('cells', []):
        ws = wb[op['sheet']]
        cell = ws.cell(row_of(ws, op['id']), headers(ws).index(op['column']) + 1)
        cell.value = text(op['new'])
        cell.data_type = 's'
    for op in patch.get('rows', []):
        insert(wb[op['sheet']], op['after'], op['values'])
    for title in {op['sheet'] for op in patch.get('rows', [])}:
        ws = wb[title]
        last = last_row(ws)
        if ws.auto_filter.ref:
            ws.auto_filter.ref = extend(ws.auto_filter.ref, last)
        for table in ws.tables.values():
            table.ref = extend(table.ref, last)
            if table.autoFilter is not None and table.autoFilter.ref:
                table.autoFilter.ref = extend(table.autoFilter.ref, last)


def refuse(problems):
    for p in problems:
        print(f'blueprint_apply: {p}', file=sys.stderr)
    print(f'blueprint_apply: {len(problems)} problem(s), nothing written', file=sys.stderr)
    return 1


def main(argv):
    parser = argparse.ArgumentParser(prog='blueprint_apply.py', description=__doc__.splitlines()[0])
    parser.add_argument('patch')
    parser.add_argument('--workbook', default=WORKBOOK)
    parser.add_argument('--out', default=OUT)
    args = parser.parse_args(argv)
    try:
        patch = json.loads(Path(args.patch).read_text(encoding='utf-8-sig'))
    except (OSError, ValueError) as e:
        return refuse([f'cannot read the patch {args.patch}: {e}'])
    workbook = Path(args.workbook)
    try:
        wb = openpyxl.load_workbook(workbook)
    except (OSError, InvalidFileException, zipfile.BadZipFile) as e:
        return refuse([f'cannot open the workbook {workbook}: {e}'])
    problems = check(wb, patch)
    if problems:
        return refuse(problems)
    apply(wb, patch)
    # Saved and exported next to the workbook first (same drive, so os.replace can move it on Windows).
    tmp = Path(tempfile.mkdtemp(prefix='.blueprint_apply-', dir=workbook.resolve().parent))
    try:
        saved, exported = tmp / workbook.name, tmp / 'blueprint'
        wb.save(saved)
        if export_blueprint.main(str(saved), str(exported)):
            print('blueprint_apply: the export refused the patched workbook, nothing written', file=sys.stderr)
            return 1
        try:
            os.replace(saved, workbook)
        except PermissionError:
            return refuse([f'cannot replace {workbook}: close it in Excel and run again'])
        try:
            shutil.copytree(exported, args.out, dirs_exist_ok=True)
        except OSError as e:
            print(f'blueprint_apply: {workbook} is patched, but the Markdown could not be written to {args.out} ({e}); '
                  f'do not apply the patch again, run: python tools/export_blueprint.py {workbook} {args.out}',
                  file=sys.stderr)
            return 1
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print(f'{workbook}: {len(patch.get("rows", []))} rows added, {len(patch.get("cells", []))} cells changed; '
          f'exported to {args.out}')
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
