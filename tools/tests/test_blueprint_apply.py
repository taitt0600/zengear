"""Tests of tools/blueprint_apply.py: a patch is applied only when every operation fits the workbook, new rows look
like the rows above them, and a refused patch leaves the workbook byte-identical and exports nothing.

Every test works on a workbook it builds in a temporary folder, or on a copy of the real one: the tool never runs
on docs/product-domain-blueprint-MVP.xlsx itself.
Needs openpyxl (CI installs it in the tools workflow); skipped without it.
Run from the repository root: python3 -m unittest discover -s tools/tests -v
"""
import json
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    openpyxl = None

REPO = Path(__file__).resolve().parents[2]
SCRIPT = REPO / 'tools' / 'blueprint_apply.py'
EXPORT = REPO / 'tools' / 'export_blueprint.py'
REAL = REPO / 'docs' / 'product-domain-blueprint-MVP.xlsx'
DECISION_HEAD = ['Decision ID', 'Topic', 'Context', 'Decision', 'Consequences', 'Source', 'Phase', 'Status']


def decision(identifier, status='Accepted'):
    return [identifier, 'Topic', 'Context', 'Decision', 'Consequences', 'Source', 'Phase 1', status]


def decision_values(identifier, status='Accepted'):
    return dict(zip(DECISION_HEAD, decision(identifier, status)))


def build(path):
    """Features (styled, with row heights), Business Rules and Decisions (with an autofilter)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Features'
    ws.append(['Feature ID', 'Feature', 'Acceptance Criteria'])
    ws.append(['AUT-04', 'Login', 'Signs in'])
    ws.append(['AUT-05', 'Logout', 'Signs out'])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(italic=True)
            cell.fill = PatternFill('solid', fgColor='FFEECC')
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 45
    ws = wb.create_sheet('Business Rules')
    ws.append(['Rule ID', 'Rule'])
    ws.append(['R-SEC-01', 'First'])
    ws.append(['R-SEC-02', 'Second'])
    ws = wb.create_sheet('Decisions')
    ws.append(DECISION_HEAD)
    ws.append(decision('DEC-001'))
    ws.append(decision('DEC-002'))
    ws.auto_filter.ref = 'A1:H3'
    wb.save(path)


@unittest.skipUnless(openpyxl, 'openpyxl is needed')
class ApplyTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.workbook = self.tmp / 'book' / 'workbook.xlsx'
        self.workbook.parent.mkdir()
        self.out = self.tmp / 'blueprint'
        build(self.workbook)

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def apply(self, patch):
        file = self.tmp / 'patch.json'
        file.write_text(patch if isinstance(patch, str) else json.dumps(patch), encoding='utf-8')
        return subprocess.run([sys.executable, str(SCRIPT), str(file), '--workbook', str(self.workbook),
                               '--out', str(self.out)], capture_output=True, text=True, encoding='utf-8')

    def applied(self, patch):
        run = self.apply(patch)
        self.assertEqual(run.returncode, 0, run.stderr)
        return openpyxl.load_workbook(self.workbook)

    def lines(self, name):
        return (self.out / name).read_text(encoding='utf-8').splitlines()

    def test_an_empty_patch_exports_the_same_markdown(self):
        reference = self.tmp / 'reference'
        reference.mkdir()
        shutil.copy(self.workbook, reference / self.workbook.name)
        run = subprocess.run([sys.executable, str(EXPORT), str(reference / self.workbook.name), str(reference / 'out')],
                             capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(run.returncode, 0, run.stderr)
        self.applied({})
        expected = sorted((reference / 'out').iterdir(), key=lambda p: p.name)
        self.assertEqual([p.name for p in expected], sorted(p.name for p in self.out.iterdir()))
        for file in expected:
            with self.subTest(file=file.name):
                self.assertEqual(file.read_bytes(), (self.out / file.name).read_bytes())

    def test_inserts_after_an_id_after_a_new_row_and_at_the_end(self):
        self.applied({'rows': [
            {'sheet': 'Business Rules', 'after': 'R-SEC-01', 'values': {'Rule ID': 'R-SEC-03', 'Rule': 'Third'}},
            {'sheet': 'Business Rules', 'after': 'R-SEC-03', 'values': {'Rule ID': 'R-SEC-04', 'Rule': 'Fourth'}},
            {'sheet': 'Decisions', 'after': 'end', 'values': decision_values('DEC-003')},
        ]})
        ids = [line.split(' | ')[0] for line in self.lines('business-rules.md')[4:]]
        self.assertEqual(ids, ['| R-SEC-01', '| R-SEC-03', '| R-SEC-04', '| R-SEC-02'])
        self.assertTrue(self.lines('decisions.md')[-1].startswith('| DEC-003 | Topic |'))

    def test_changes_a_cell(self):
        wb = self.applied({'cells': [{'sheet': 'Features', 'id': 'AUT-04', 'column': 'Feature', 'old': 'Login',
                                      'new': 'Sign in'}]})
        self.assertEqual(wb['Features']['B2'].value, 'Sign in')
        self.assertIn('| AUT-04 | Sign in | Signs in |', self.lines('features.md'))

    def test_a_newline_is_a_line_break_and_a_dash_stays_text(self):
        wb = self.applied({'rows': [{'sheet': 'Features', 'after': 'AUT-04', 'values': {
            'Feature ID': 'AUT-06', 'Feature': '-1 dash', 'Acceptance Criteria': 'first\r\nsecond'}}]})
        self.assertIn('| AUT-06 | -1 dash | first<br>second |', self.lines('features.md'))
        cell = wb['Features']['B3']
        self.assertEqual((cell.value, cell.data_type), ('-1 dash', 's'))
        self.assertEqual(wb['Features']['C3'].value, 'first\nsecond')

    def test_a_new_row_takes_the_style_and_height_of_the_row_above(self):
        ws = self.applied({'rows': [{'sheet': 'Features', 'after': 'AUT-04', 'values': {
            'Feature ID': 'AUT-06', 'Feature': 'Profile'}}]})['Features']
        self.assertEqual([c.value for c in ws['A']], ['Feature ID', 'AUT-04', 'AUT-06', 'AUT-05'])
        for column in 'ABC':
            with self.subTest(column=column):
                new, above = ws[f'{column}3'], ws[f'{column}2']
                self.assertTrue(new.font.italic)
                self.assertEqual(new.fill.fgColor.rgb, above.fill.fgColor.rgb)
                self.assertTrue(new.alignment.wrap_text)
        self.assertEqual(ws.row_dimensions[3].height, 30)
        self.assertEqual(ws.row_dimensions[4].height, 45, 'the row below keeps its own height')

    def test_extends_the_autofilter_to_the_last_row(self):
        wb = self.applied({'rows': [{'sheet': 'Decisions', 'after': 'end', 'values': decision_values('DEC-003')}]})
        self.assertEqual(wb['Decisions'].auto_filter.ref, 'A1:H4')

    def assert_refused(self, patch, *fragments):
        before = self.workbook.read_bytes()
        run = self.apply(patch)
        self.assertEqual(run.returncode, 1, run.stdout + run.stderr)
        for fragment in fragments:
            self.assertIn(fragment, run.stderr)
        self.assertIn('nothing written', run.stderr)
        self.assertEqual(self.workbook.read_bytes(), before, 'the workbook must not change')
        self.assertFalse(self.out.exists(), 'nothing may be exported')
        self.assertEqual(list(self.workbook.parent.iterdir()), [self.workbook], 'no temporary folder may be left')

    def test_refusals(self):
        cell = {'sheet': 'Features', 'id': 'AUT-04', 'column': 'Feature', 'old': 'Login', 'new': 'Sign in'}
        row = {'sheet': 'Decisions', 'after': 'end', 'values': decision_values('DEC-003')}
        cases = {
            'existing ID': ({'rows': [{**row, 'values': decision_values('DEC-002')}]}, 'DEC-002 is already in Decisions'),
            'wrong old': ({'cells': [{**cell, 'old': 'Log in'}]}, "the patch expects 'Log in'"),
            'unknown sheet': ({'rows': [{**row, 'sheet': 'Decision'}]}, "no sheet 'Decision'"),
            'unknown column': ({'cells': [{**cell, 'column': 'Name'}]}, "no column 'Name' in Features"),
            'unknown after': ({'rows': [{**row, 'after': 'DEC-009'}]}, "'DEC-009', is not in Decisions"),
            'formula': ({'cells': [{**cell, 'new': '=1+1'}]}, 'starts with ='),
            'not text': ({'rows': [{**row, 'values': {**decision_values('DEC-003'), 'Phase': 1}}]}, 'Phase is not text'),
            'hidden character': ({'cells': [{**cell, 'new': 'Sign​in'}]}, 'U+200B'),
            'ID column': ({'cells': [{**cell, 'column': 'Feature ID', 'old': 'AUT-04', 'new': 'AUT-07'}]},
                          'is the ID column'),
            'unknown key': ({'row': [row]}, "unknown key 'row'"),
            'unknown column of a row': ({'rows': [{**row, 'values': {**decision_values('DEC-003'), 'Owner': 'x'}}]},
                                        "no column 'Owner' in Decisions"),
            'no ID': ({'rows': [{**row, 'values': {'Topic': 'x'}}]}, "no value for the ID column 'Decision ID'"),
            'same cell twice': ({'cells': [cell, {**cell, 'new': 'Log on'}]}, 'is changed twice'),
            'same after twice': ({'rows': [{**row, 'after': 'DEC-001'},
                                           {**row, 'after': 'DEC-001', 'values': decision_values('DEC-004')}]},
                                 'another row is already inserted after DEC-001'),
            'bad JSON': ('{"rows": [', 'cannot read the patch'),
            'refused by the export': ({'rows': [{**row, 'values': decision_values('DEC-003', 'Accepted ')}]},
                                      "the Status 'Accepted '"),
        }
        for name, (patch, fragment) in cases.items():
            with self.subTest(case=name):
                self.assert_refused(patch, fragment)

    def test_lists_every_problem(self):
        self.assert_refused({'rows': [{'sheet': 'Nope', 'after': 'end', 'values': {}}],
                             'cells': [{'sheet': 'Features', 'id': 'AUT-09', 'column': 'Feature', 'old': '', 'new': 'x'}]},
                            "no sheet 'Nope'", "'AUT-09' is not in Features", '2 problem(s)')

    def test_a_missing_workbook_is_refused(self):
        self.workbook.unlink()
        run = self.apply({})
        self.assertEqual(run.returncode, 1, run.stderr)
        self.assertIn('cannot open the workbook', run.stderr)
        self.assertFalse(self.out.exists())

    def test_the_same_patch_twice_is_refused(self):
        patch = {'rows': [{'sheet': 'Decisions', 'after': 'end', 'values': decision_values('DEC-003')}],
                 'cells': [{'sheet': 'Features', 'id': 'AUT-04', 'column': 'Feature', 'old': 'Login', 'new': 'Sign in'}]}
        self.applied(patch)
        shutil.rmtree(self.out)
        self.assert_refused(patch, 'DEC-003 is already in Decisions', "AUT-04 Feature is 'Sign in'")


@unittest.skipUnless(openpyxl, 'openpyxl is needed')
class RealWorkbookCopyTests(unittest.TestCase):
    """On a copy of the real workbook: a save by openpyxl changes nothing the export shows."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.workbook = self.tmp / 'book' / REAL.name
        self.workbook.parent.mkdir()
        shutil.copy(REAL, self.workbook)
        self.out = self.tmp / 'blueprint'
        self.committed = REPO / 'docs' / 'blueprint'

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def apply(self, patch):
        file = self.tmp / 'patch.json'
        file.write_text(json.dumps(patch), encoding='utf-8')
        run = subprocess.run([sys.executable, str(SCRIPT), str(file), '--workbook', str(self.workbook),
                              '--out', str(self.out)], capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(run.returncode, 0, run.stderr)

    def lines(self, folder, name):
        return (folder / name).read_text(encoding='utf-8').splitlines()

    def test_an_empty_patch_exports_docs_blueprint(self):
        self.apply({})
        self.assertEqual(sorted(p.name for p in self.out.iterdir()), sorted(p.name for p in self.committed.iterdir()))
        for file in self.committed.iterdir():
            with self.subTest(file=file.name):
                self.assertEqual(self.lines(self.out, file.name), self.lines(self.committed, file.name))

    def test_one_new_decision_adds_only_its_row(self):
        head = [h.strip() for h in self.lines(self.committed, 'decisions.md')[2].strip('|').split(' | ')]
        values = {h: f'Test {h}' for h in head}
        values.update({head[0]: 'DEC-999', 'Status': 'Proposed'})
        self.apply({'rows': [{'sheet': 'Decisions', 'after': 'end', 'values': values}]})
        for file in self.committed.iterdir():
            with self.subTest(file=file.name):
                exported, committed = self.lines(self.out, file.name), self.lines(self.committed, file.name)
                if file.name == 'decisions.md':
                    self.assertEqual(exported[:-1], committed)
                    self.assertTrue(exported[-1].startswith('| DEC-999 | Test '), exported[-1])
                elif file.name == 'README.md':
                    count = lambda m: f'— {int(m.group(1)) + 1} rows'
                    self.assertEqual(exported, [re.sub(r'— (\d+) rows', count, line)
                                                if line.startswith('- [Decisions]') else line for line in committed])
                else:
                    self.assertEqual(exported, committed)


if __name__ == '__main__':
    unittest.main()
